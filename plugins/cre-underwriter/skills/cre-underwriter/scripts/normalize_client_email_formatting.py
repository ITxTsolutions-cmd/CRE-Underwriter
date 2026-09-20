#!/usr/bin/env python3
"""Normalize CRE Underwriter client-email DOCX typography without changing layout.

Usage:
    python normalize_client_email_formatting.py input.docx [output.docx]

If output.docx is omitted, the input file is updated in place.
"""
from pathlib import Path
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT = "Georgia"
SECTION_HEADINGS = {
    "Transaction Summary",
    "Key Issues",
    "Risk Factors, Exceptions & Mitigants",
    "Inconsistencies to Resolve",
    "Additional Information & Documentation Needed",
}
TABLE_HEADINGS = {
    "Preliminary Soft Quote / Indicative Terms",
    "Loan Terms",
    "Indicative Structure",
}

def set_rfonts(rpr, font_name=FONT):
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)

def set_run_font(run, font_name=FONT):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    set_rfonts(rpr, font_name)

def iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)

def set_document_defaults(doc):
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, doc_defaults)
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        doc_defaults.append(rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    set_rfonts(rpr, FONT)

    for style in doc.styles:
        try:
            style.font.name = FONT
            rpr = style.element.get_or_add_rPr()
            set_rfonts(rpr, FONT)
        except Exception:
            pass

def normalize(doc):
    set_document_defaults(doc)

    for paragraph in iter_paragraphs(doc):
        for run in paragraph.runs:
            set_run_font(run)

    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in iter_paragraphs(part):
                for run in paragraph.runs:
                    set_run_font(run)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip().rstrip(".")
        if text in SECTION_HEADINGS:
            for run in paragraph.runs:
                run.bold = True
                set_run_font(run)

    ts_index = None
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().rstrip(".") == "Transaction Summary":
            ts_index = i
            break
    if ts_index is not None:
        nonempty = [i for i in range(ts_index) if doc.paragraphs[i].text.strip()]
        for i in nonempty:
            for run in doc.paragraphs[i].runs:
                run.bold = False
                set_run_font(run)

    nonempty_body = [p for p in doc.paragraphs if p.text.strip()]
    if nonempty_body:
        last = nonempty_body[-1]
        if last.text.strip().rstrip(".") not in SECTION_HEADINGS:
            for run in last.runs:
                run.bold = False
                set_run_font(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = " ".join(p.text for p in cell.paragraphs).strip().rstrip(".")
                if text in TABLE_HEADINGS:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                            set_run_font(run)

def validate(doc):
    errors = []
    for paragraph in iter_paragraphs(doc):
        for run in paragraph.runs:
            if run.text and run.font.name not in (FONT, None):
                errors.append(f"Non-Georgia run: {run.text[:40]!r} -> {run.font.name!r}")
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip().rstrip(".")
        if text in SECTION_HEADINGS and any(run.bold is not True for run in paragraph.runs if run.text):
            errors.append(f"Section heading not fully bold: {text}")
    return errors

def main():
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Usage: normalize_client_email_formatting.py input.docx [output.docx]")
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) == 3 else src
    doc = Document(src)
    normalize(doc)
    doc.save(dst)
    check = Document(dst)
    errors = validate(check)
    if errors:
        raise SystemExit("Formatting validation failed:\n" + "\n".join(errors))
    print(dst)

if __name__ == "__main__":
    main()
